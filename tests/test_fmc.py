"""FMC abstract: parser, conference limits, DOCX build (synthetic text)."""
import pytest

from vnsoc import fmc

SECT_VI = ["TIÊU ĐỀ", "ĐẶT VẤN ĐỀ", "MỤC TIÊU", "PHƯƠNG PHÁP NGHIÊN CỨU", "KẾT QUẢ", "KẾT LUẬN", "TỪ KHÓA"]
SECT_EN = ["TITLE", "BACKGROUND", "OBJECTIVE", "METHODS", "RESULTS", "CONCLUSION", "KEYWORDS"]


def md(box="Tóm tắt ngắn.", body="nội dung ngắn gọn", kw_vi="a; b; c", kw_en="x; y; z", title="Tên đề tài"):
    out = ["# x", "", "## Ô TÓM TẮT", "", box, "", "## VI", ""]
    for s in SECT_VI:
        out += [f"### {s}", "", title if s == "TIÊU ĐỀ" else kw_vi if s == "TỪ KHÓA" else body, ""]
    out += ["## EN", ""]
    for s in SECT_EN:
        out += [f"### {s}", "", title if s == "TITLE" else kw_en if s == "KEYWORDS" else body, ""]
    return "\n".join(out + ["## GHI CHÚ", "", "Chưa đăng.", "", "## NOTE", "", "Not published.", "",
                           "## ABSTRACT BOX", "", "Short abstract.", ""])


def test_parse_and_ok():
    d = fmc.parse(md())
    assert d["box"] == "Tóm tắt ngắn." and list(d["vi"])[0] == "TIÊU ĐỀ" and d["note"] == "Chưa đăng."
    assert d["box_en"] == "Short abstract." and d["note_en"] == "Not published."
    assert fmc.checks(d) == [] and fmc.warnings(d) == []


@pytest.mark.parametrize("kw,problem", [
    (dict(box="x" * 501), "ký tự"), (dict(body=" ".join(["từ"] * 110)), "từ >"), (dict(kw_vi="a; b"), "từ khóa"),
    (dict(kw_en="X; y; z"), "từ khóa"), (dict(title="t" * 151), "tiêu đề"), (dict(body="có {{design.x}}"), "placeholder"),
])
def test_limits(kw, problem):
    assert any(problem in p for p in fmc.checks(fmc.parse(md(**kw))))


def test_recommended_length_is_a_warning_only():
    d = fmc.parse(md(body=" ".join(["từ"] * 60)))            # 5 sections x 60 = 300 words: > 250, <= 500
    assert fmc.checks(d) == [] and any("khuyến khích" in w for w in fmc.warnings(d))


def test_docx_one_language_per_file(tmp_path):
    pytest.importorskip("docx")
    import docx

    au = [{"name": "Binh Minh Ngo", "affiliation": "Khoa CNTT, TDTU, Việt Nam", "affiliation_en": "FIT, TDTU, Vietnam",
           "email": "a@b.vn", "corresponding": True},
          {"name": "B", "affiliation": "Khoa Y, PCTU, Việt Nam", "affiliation_en": "Faculty of Medicine, PCTU, Vietnam"}]
    d = fmc.parse(md())
    for lang in ("vi", "en"):
        out = tmp_path / f"a_{lang}.docx"
        fmc.build_docx(d, out, au, lang)
        text = "\n".join(p.text for p in docx.Document(out).paragraphs)
        assert "TÊN ĐỀ TÀI" in text and "Binh Minh Ngo¹, B²" in text and 0 < out.stat().st_size < 1_000_000
        if lang == "vi":
            assert "ĐẶT VẤN ĐỀ: nội dung" in text and "KEYWORDS" not in text and "Chưa đăng." in text
            assert "²Khoa Y, PCTU, Việt Nam" in text and "Tác giả liên hệ: Binh Minh Ngo, a@b.vn" in text
        else:
            assert "KEYWORDS: x; y; z" in text and "ĐẶT VẤN ĐỀ" not in text and "Not published." in text
            assert "²Faculty of Medicine, PCTU, Vietnam" in text and "Corresponding author: Binh Minh Ngo" in text
            assert "Khoa" not in text


def test_docx_from_conference_template(tmp_path):
    pytest.importorskip("docx")
    import docx

    from vnsoc.paths import paths

    tpl = paths().root / fmc.TEMPLATE
    if not tpl.exists():
        pytest.skip("mẫu hội nghị không có trong repo")
    au = [{"name": "Binh Minh Ngo", "name_vi": "Ngô Bình Minh", "affiliation": "Khoa CNTT, TDTU, Việt Nam",
           "affiliation_en": "FIT, TDTU, Vietnam", "email": "a@b.vn", "corresponding": True},
          {"name": "B", "affiliation": "Khoa Y, PCTU, Việt Nam", "affiliation_en": "Faculty of Medicine, PCTU, Vietnam"},
          {"name": "C", "affiliation": "Khoa RHM, PCTU, Việt Nam", "affiliation_en": "Faculty of Dentistry, PCTU, Vietnam"}]
    d = fmc.parse(md().replace("Chưa đăng.", "HÌNH THỨC BÁO CÁO: Oral\nGHI CHÚ: Chưa đăng.")
                  .replace("Not published.", "PRESENTATION FORMAT: Oral\nNOTE: Not published."))
    assert d["note"] == "HÌNH THỨC BÁO CÁO: Oral\nGHI CHÚ: Chưa đăng."
    for lang in ("vi", "en"):
        out = tmp_path / f"t_{lang}.docx"
        fmc.build_docx_template(d, out, au, lang, tpl)
        doc = docx.Document(out)
        text = "\n".join(p.text for p in doc.paragraphs)
        for gone in ("MẪU", "Điền nội dung", "[Tiêu đề", "[Tên của", "HẠN NỘP", "Tên đầy đủ tác giả", "Khoa/ Bộ môn",
                     "Công nghệ Gen"):
            assert gone not in text
        assert "TÊN ĐỀ TÀI" in text and len(doc.sections[0].header._element.xpath(".//a:blip")) == 2   # logos kept
        au_p = next(p for p in doc.paragraphs if p.text.startswith(("Ngô Bình Minh", "Binh Minh Ngo")))
        assert [r.text for r in au_p.runs if r.font.superscript] == ["1", "2", "3"]
        if lang == "vi":
            assert "Ngô Bình Minh1, B2, C3" in text and "ĐẶT VẤN ĐỀ: nội dung ngắn gọn" in text
            assert "3Khoa RHM, PCTU, Việt Nam" in text and "HÌNH THỨC BÁO CÁO: Oral" in text and "KEYWORDS" not in text
        else:
            assert "Binh Minh Ngo1" in text and "BACKGROUND: nội dung" in text and "KEYWORDS: x; y; z" in text
            assert "ĐẶT VẤN ĐỀ" not in text and "Khoa" not in text and "HỘI NGHỊ" not in text
            assert "PRESENTATION FORMAT: Oral" in text and "Corresponding author: Binh Minh Ngo, a@b.vn" in text


def test_nbsp_keeps_units_together():
    assert fmc.nbsp("của Bộ Y tế (p = 0,15)") == "của Bộ\u00a0Y\u00a0tế (p\u00a0=\u00a00,15)"
