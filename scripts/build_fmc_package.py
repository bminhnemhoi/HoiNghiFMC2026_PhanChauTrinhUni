"""Build the FMC 2026 submission folder (Nop_Final_PhanChauTrinh_HT2026/) from the rendered abstracts.

Inputs (render first: $PY -m vnsoc.numbers render):
  manuscript/build/fmc/abstract_fmc.md      main attached abstract (organisers: ≤ 500 words, one language per file)
  manuscript/build/fmc/abstract_fmc_250.md  backup at the template's recommended 250 words
Outputs:
  01_Abstract_TiengViet.docx/.pdf, 02_Abstract_TiengAnh.docx/.pdf   conference template filled (vnsoc.fmc)
  03_NOI_DUNG_DIEN_FORM.txt   plain text for each web-form field, checked against the form's real limits
  00_HUONG_DAN_NOP.pdf        formatted guide: steps, field table, quality checklist, SHA-256 of the files
  Du_phong_ban_250_tu/        the 250-word version (DOCX + PDF)
Web-form limits were read from the live form (conference.pctu.edu.vn/dang-ki/huong-dan-dang-ki-de-tai/, 27/9/2026):
Tendetai maxlength 150, Donvi 400, tacgia 400, tomtat 500 (character counter), file ≤ 1 MB (jpg/png/txt/doc/docx/pdf).
PDFs are exported by Microsoft Word (COM) so they look exactly like the DOCX.

  $PY scripts/build_fmc_package.py
"""
from __future__ import annotations

import hashlib
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from vnsoc import fmc  # noqa: E402

OUT = ROOT / "Nop_Final_PhanChauTrinh_HT2026"
BUILD = ROOT / "manuscript" / "build" / "fmc"
FORM_URL = "conference.pctu.edu.vn → Đăng kí → Hướng dẫn đăng kí đề tài (biểu mẫu nộp đề tài)"
LIMITS = {"Tendetai": 150, "Donvi": 400, "tacgia": 400, "tomtat": 500}
TOPIC = "1. AI và chuyển đổi số trong Y tế và Giáo dục Y khoa"
FILE_MAX = 1_000_000

PS_EXPORT = r"""param([string]$src, [string]$dst)
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try { $d = $w.Documents.Open($src, $false, $true); $d.SaveAs2($dst, 17); $d.Close($false) } finally { $w.Quit() }
"""


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def to_pdf(docx: Path) -> Path:
    pdf = docx.with_suffix(".pdf")
    with tempfile.TemporaryDirectory() as t:
        ps = Path(t) / "export.ps1"
        ps.write_text(PS_EXPORT, encoding="utf-8")
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps),
                        "-src", str(docx), "-dst", str(pdf)], check=True, capture_output=True)
    return pdf


def build_abstracts(stem: str, dest: Path) -> dict:
    src = BUILD / f"{stem}.md"
    doc = fmc.parse(src.read_text(encoding="utf-8"))
    probs = fmc.checks(doc)
    if probs:
        raise SystemExit(f"{src.name}: {probs}")
    cfg = yaml.safe_load((ROOT / "configs" / "project.yaml").read_text(encoding="utf-8"))
    tpl = ROOT / fmc.TEMPLATE
    dest.mkdir(parents=True, exist_ok=True)
    files = {}
    for lang, name in (("vi", "01_Abstract_TiengViet"), ("en", "02_Abstract_TiengAnh")):
        out = dest / f"{name}.docx"
        fmc.build_docx_template(doc, out, cfg["authors"], lang, tpl)
        if out.stat().st_size > FILE_MAX:
            raise SystemExit(f"{out.name} > 1 MB")
        files[lang] = (out, to_pdf(out))
    return {"doc": doc, "files": files, "authors": cfg["authors"]}


def form_fields(doc: dict, authors: list[dict]) -> list[tuple[str, str, str, int | None]]:
    """(form label, field name, text, limit) in the order of the web form."""
    aff: list[str] = []
    for a in authors:
        if a["affiliation"] not in aff:
            aff.append(a["affiliation"])
    donvi = "; ".join(f"({k + 1}) {x}" for k, x in enumerate(aff))
    tacgia = "; ".join(f"{a.get('name_vi') or a['name']} ({aff.index(a['affiliation']) + 1})"
                       + (" – tác giả liên hệ" if a.get("corresponding") else "") for a in authors)
    corr = next(a for a in authors if a.get("corresponding"))
    title = list(doc["vi"].items())[0][1]
    rows = [("Chọn Chủ đề", "chude[]", f"Tích ô: {TOPIC}", None),
            ("Tên đề tài", "Tendetai", nfc(title), LIMITS["Tendetai"]),
            ("Đơn vị", "Donvi", nfc(donvi), LIMITS["Donvi"]),
            ("Tác giả / Đồng tác giả / Tác giả liên hệ", "tacgia", nfc(tacgia), LIMITS["tacgia"]),
            ("Tóm tắt đề tài / Abstract", "tomtat", nfc(doc["box"]), LIMITS["tomtat"]),
            ("Email liên hệ", "email", corr["email"], None),
            ("File đính kèm", "file", "01_Abstract_TiengViet.docx (hoặc .pdf)", None)]
    for label, _, text, lim in rows:
        if lim and len(text) > lim:
            raise SystemExit(f"ô '{label}': {len(text)} ký tự > {lim}")
    return rows


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def write_form_txt(rows, doc: dict, dest: Path) -> None:
    en_title = list(doc["en"].items())[0][1]
    lines = ["NỘI DUNG ĐIỀN VÀO BIỂU MẪU NỘP ĐỀ TÀI — HỘI NGHỊ KHOA HỌC FMC 2026",
             f"Biểu mẫu: {FORM_URL}",
             "Chép NGUYÊN VĂN từng khối vào đúng ô (không thêm bớt). Giới hạn ký tự là giới hạn thật của biểu mẫu (27/9/2026).",
             "=" * 100, ""]
    for label, _, text, lim in rows:
        head = f"[{label}]" + (f"  —  {len(text)}/{lim} ký tự" if lim else "")
        lines += [head, text, ""]
    lines += ["-" * 100, "BẢN TIẾNG ANH (chỉ dùng nếu Ban tổ chức yêu cầu tiếng Anh; khi đó đính kèm 02_Abstract_TiengAnh.docx)", "",
              f"[Title]  —  {len(en_title)}/150", en_title, "",
              f"[Abstract field]  —  {len(nfc(doc['box_en']))}/500", nfc(doc["box_en"]), ""]
    (dest / "03_NOI_DUNG_DIEN_FORM.txt").write_text("\n".join(lines), encoding="utf-8")


def write_guide_pdf(rows, main: dict, backup: dict, dest: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    BLUE = RGBColor(0x15, 0x3D, 0x63)
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(3)

    def para(text="", bold=False, size=11, color=None, align=None, italic=False):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)
        if color:
            r.font.color.rgb = color
        if align is not None:
            p.alignment = align
        return p

    def head(text):
        p = para(text, bold=True, size=13, color=BLUE)
        p.paragraph_format.space_before = Pt(10)

    def shade(cell, hex_fill):
        tc = cell._element.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"), sh.set(qn("w:color"), "auto"), sh.set(qn("w:fill"), hex_fill)
        tc.append(sh)

    def table(header, body, widths):
        t = d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size = True, Pt(10)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, "153D63")
        for row in body:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(10)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        return t

    vi = main["doc"]["vi"]
    para("HỘI NGHỊ KHOA HỌC FMC 2026", bold=True, size=16, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Trường Đại học Phan Châu Trinh — conference.pctu.edu.vn", size=11, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("HƯỚNG DẪN NỘP TÓM TẮT BÁO CÁO", bold=True, size=13, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(nfc(list(vi.items())[0][1]), italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Hạn nộp: hết ngày 30/9/2026 (Ban tổ chức xác nhận nộp trong ngày 30/9 vẫn được nhận).", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER)

    head("1. Các file trong thư mục")
    f = main["files"]
    table(["File", "Dùng để", "Dung lượng"], [
        ("01_Abstract_TiengViet.docx", "FILE ĐÍNH KÈM CHÍNH — tóm tắt tiếng Việt điền đúng mẫu của hội nghị", f"{f['vi'][0].stat().st_size // 1024} KB"),
        ("01_Abstract_TiengViet.pdf", "Bản PDF của file trên (nếu muốn giữ nguyên trình bày)", f"{f['vi'][1].stat().st_size // 1024} KB"),
        ("02_Abstract_TiengAnh.docx / .pdf", "Bản tiếng Anh (chỉ khi Ban tổ chức yêu cầu tiếng Anh)", f"{f['en'][0].stat().st_size // 1024} KB"),
        ("03_NOI_DUNG_DIEN_FORM.txt", "Chữ để chép vào từng ô của biểu mẫu web (dạng chữ thuần, không lỗi định dạng)", "—"),
        ("Du_phong_ban_250_tu/", "Bản rút gọn 250 từ (mức khuyến khích của mẫu) — chỉ dùng nếu muốn nộp bản ngắn", "—"),
    ], [5.2, 9.3, 2.2])

    head("2. Các bước nộp (bạn tự thực hiện)")
    for k, s in enumerate([
        "Mở biểu mẫu: " + FORM_URL + ".",
        "Điền từng ô theo bảng ở mục 3 — chép chữ từ file 03_NOI_DUNG_DIEN_FORM.txt.",
        "Ô File đính kèm: chọn 01_Abstract_TiengViet.docx (hoặc .pdf).",
        "Bấm Gửi; chụp màn hình trang/email xác nhận, ghi lại mã hồ sơ nếu có.",
        "Báo lại trong Claude Code bằng một dòng bắt đầu: XONG HG1.9 đã nộp lúc <giờ, ngày>; mã: <nếu có>.",
    ], 1):
        para(f"Bước {k}. {s}")

    head("3. Nội dung từng ô của biểu mẫu (giới hạn thật của biểu mẫu)")
    table(["Ô trên biểu mẫu", "Nội dung", "Độ dài / giới hạn"],
          [(lab, txt, f"{len(txt)}/{lim} ký tự" if lim else "—") for lab, _, txt, lim in rows], [3.6, 10.8, 2.3])
    para("Lưu ý: ô “Tóm tắt đề tài / Abstract” của biểu mẫu đếm KÝ TỰ và chặn ở 500 ký tự (maxlength=500, bộ đếm "
         "ký tự), nên chỉ điền đoạn tóm tắt ngắn; tóm tắt đầy đủ (giới hạn 500 TỪ theo Ban tổ chức) nằm trong file đính kèm.",
         italic=True, size=10)

    head("4. Kiểm tra chất lượng đã thực hiện")
    wv, we = fmc.body_words(main["doc"], "vi"), fmc.body_words(main["doc"], "en")
    bv, be = fmc.body_words(backup["doc"], "vi"), fmc.body_words(backup["doc"], "en")
    checks = [
        "Điền trực tiếp vào file mẫu chính thức của hội nghị: logo, tiêu đề, màu, phông Times New Roman, tiêu đề in hoa cỡ 13, nội dung cỡ 12, giãn dòng 1,5.",
        "Đủ các mục của mẫu: Đặt vấn đề, Mục tiêu, Phương pháp nghiên cứu, Kết quả, Kết luận, Từ khóa (5 từ khóa, viết thường, cách nhau bằng dấu chấm phẩy); cuối bài ghi hình thức báo cáo, lĩnh vực, tình trạng công bố.",
        f"Độ dài (năm mục, đếm cả tên mục): bản chính {wv} từ (tiếng Việt), {we} từ (tiếng Anh) — trong giới hạn 500 từ; bản dự phòng {bv}/{be} từ.",
        "Mỗi file một ngôn ngữ, theo yêu cầu của Ban tổ chức; file < 1 MB, định dạng DOCX/PDF được chấp nhận.",
        "Mọi con số lấy tự động từ kho số liệu do mã phân tích ghi (không gõ tay); kiểm toán độc lập đã tính lại các số từ dữ liệu gốc và chạy lại bộ chấm (482/482 nhãn khớp).",
        "Giới hạn nêu rõ ở cả hai ngôn ngữ: thí điểm, mẫu chọn chủ đích, kiểm tra do AI thực hiện, chưa có bác sĩ duyệt; kiểm định McNemar ghi là phân tích khám phá.",
    ]
    for c in checks:
        para("☑ " + c, size=10.5)

    head("5. Mã băm SHA-256 (đối chiếu đúng file đã nộp)")
    hs = [(p.name, sha(p)) for lang in ("vi", "en") for p in main["files"][lang]]
    table(["File", "SHA-256"], hs, [5.2, 11.5])
    para("Tác giả: Ngô Bình Minh (tác giả liên hệ, ngobinhminh.st@tdtu.edu.vn), Ngô Bình Thống, Trần Đoàn Mai Hương. "
         "Nhớ báo cho hai đồng tác giả trước khi nộp. Nếu muốn đổi hình thức báo cáo (Oral → Poster), báo Claude dựng lại.",
         italic=True, size=10)
    tmp = dest / "00_HUONG_DAN_NOP.docx"
    d.save(tmp)
    to_pdf(tmp)
    tmp.unlink()


def main() -> int:
    OUT.mkdir(exist_ok=True)
    main_ = build_abstracts("abstract_fmc", OUT)
    backup = build_abstracts("abstract_fmc_250", OUT / "Du_phong_ban_250_tu")
    rows = form_fields(main_["doc"], main_["authors"])
    write_form_txt(rows, main_["doc"], OUT)
    write_guide_pdf(rows, main_, backup, OUT)
    for old in ("00_DOC_TRUOC_HUONG_DAN_NOP.txt", "03_NOI_DUNG_CHEP_VAO_FORM.txt", "04_KIEM_TRA_TRUOC_KHI_NOP.txt"):
        if (OUT / old).exists():
            (OUT / old).unlink()
    print(f"OK: {OUT}")
    for p in sorted(OUT.rglob("*")):
        if p.is_file():
            print(f"  {p.relative_to(OUT)}  {p.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
